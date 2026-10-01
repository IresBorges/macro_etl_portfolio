import logging
import os

import pandas as pd
import openpyxl

from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)

from openpyxl.utils import get_column_letter

from src.config import (
    DATA_PROC_DIR,
    OUTPUT_DIR,
    LOGS_DIR
)

logger = logging.getLogger(
    "ETL_Load"
)


def generate_excel_dashboard() -> None:
    """
    Gera o dashboard Excel
    a partir dos dados transformados.
    """

    input_path = os.path.join(
        DATA_PROC_DIR,
        "macro_transformed.csv"
    )

    output_path = os.path.join(
        OUTPUT_DIR,
        "macro_dashboard.xlsx"
    )

    if not os.path.exists(input_path):

        logger.error(
            f"Ficheiro não encontrado: {input_path}"
        )

        raise FileNotFoundError(
            f"Ficheiro não encontrado: {input_path}"
        )

    logger.info(
        "A iniciar a geração do dashboard Excel."
    )

    df = pd.read_csv(
        input_path,
        parse_dates=["date"]
    )

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Macro Overview"

    ws.views.sheetView[0].showGridLines = True

    font_family = "Segoe UI"

    header_fill = PatternFill(
        start_color="1F4E78",
        end_color="1F4E78",
        fill_type="solid"
    )

    header_font = Font(
        name=font_family,
        size=11,
        bold=True,
        color="FFFFFF"
    )

    title_font = Font(
        name=font_family,
        size=16,
        bold=True,
        color="1F4E78"
    )

    data_font = Font(
        name=font_family,
        size=10
    )

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    ws["A1"] = (
        "Relatório Macroeconómico - Pipeline ETL"
    )

    ws["A1"].font = title_font

    ws.append([])

    headers = [
        "Data",
        "Inflação YoY (%)",
        "Taxa de Juros (FEDFUNDS)"
    ]

    ws.append(headers)

    for col_num in range(
        1,
        len(headers) + 1
    ):

        cell = ws.cell(
            row=3,
            column=col_num
        )

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        cell.border = thin_border

    ws.row_dimensions[3].height = 25

    for _, row in df.iterrows():

        date_str = row["date"].strftime(
            "%Y-%m-%d"
        )

        ws.append([
            date_str,
            row["CPI_Inflation_YoY"],
            row["Interest_Rate"]
        ])

    for row_num in range(
        4,
        4 + len(df)
    ):

        for col_num in range(1, 4):

            cell = ws.cell(
                row=row_num,
                column=col_num
            )

            cell.font = data_font
            cell.border = thin_border

        ws.cell(
            row=row_num,
            column=1
        ).alignment = Alignment(
            horizontal="center"
        )

        ws.cell(
            row=row_num,
            column=2
        ).number_format = "0.00"

        ws.cell(
            row=row_num,
            column=3
        ).number_format = "0.00"

    for col in ws.columns:

        max_length = max(
            len(str(cell.value or ""))
            for cell in col
        )

        col_letter = get_column_letter(
            col[0].column
        )

        ws.column_dimensions[
            col_letter
        ].width = max(
            max_length + 5,
            18
        )

    wb.save(output_path)

    logger.info(
        f"Dashboard Excel gerado com sucesso em: "
        f"{output_path}"
    )


if __name__ == "__main__":

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
        handlers=[
            logging.FileHandler(
                os.path.join(
                    LOGS_DIR,
                    "pipeline.log"
                ),
                encoding="utf-8"
            ),
            logging.StreamHandler()
        ]
    )

    generate_excel_dashboard()
          